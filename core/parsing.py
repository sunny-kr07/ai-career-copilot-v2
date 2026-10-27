import io
import re

import pdfplumber
from docx import Document

MIN_CHARS = 200


class ParseError(Exception):
    """Raised when a resume file can't be turned into usable text."""


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)        # collapse spaces/tabs
    text = re.sub(r"\n{3,}", "\n\n", text)     # collapse blank lines
    return text.strip()


def _read_pdf(data: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(page.extract_text() or "") for page in pdf.pages]
    except Exception as e:
        raise ParseError(f"Could not read this PDF: {e}")
    return "\n".join(pages)


def _read_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Could not read this DOCX: {e}")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:                    # resumes often use tables
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_text(file_bytes: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _read_pdf(file_bytes)
    elif name.endswith(".docx"):
        text = _read_docx(file_bytes)
    else:
        raise ParseError("Unsupported file type. Please upload a PDF or DOCX.")

    text = _clean(text)
    if len(text) < MIN_CHARS:
        raise ParseError(
            "Very little text was found. The file may be scanned or image-based. "
            "Try a text-based PDF or a DOCX."
        )
    return text